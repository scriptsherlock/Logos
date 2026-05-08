from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent

RUBRIC_TITLE = "Sentiment Analysis Communication Rubric"
RUBRIC_FILE_NAME = "nlp-sentiment-communication-rubric.docx"
RUBRIC_TEXT_FILE_NAME = "nlp-sentiment-communication-rubric.md"
TEACHER_CONFIG_FILE_NAME = "nlp-sentiment-teacher-config.json"
TEACHER_BUNDLE_FILE_NAME = "nlp-sentiment-teacher-bundle.json"

TASK_OVERVIEW = {
    "title": "Coursework Task: Sentiment Analysis with Error Analysis",
    "summary": "Build a sentiment classification model for movie reviews and critically evaluate its performance.",
    "requirements": [
        "Use any dataset, such as IMDb or a small custom review dataset.",
        "Train a basic model, for example logistic regression or a simple neural network.",
        "Report the methodology, evaluation metrics, and the reasons behind your design choices.",
        "Include evaluation metrics such as accuracy and F1 score.",
        "Perform error analysis and identify where the model fails.",
        "Reflect on limitations and realistic improvements.",
    ],
    "deliverable": (
        "Write a 1500-2000 word report with an Introduction, Methodology, Results, Error Analysis, and Conclusion."
    ),
}

BAND_MAPPING = [
    ("0-39", "0.0-0.39", "Disorganized, unclear, poor academic practice"),
    ("40-49", "0.4-0.49", "Basic structure, weak clarity"),
    ("50-59", "0.5-0.59", "Reasonable but inconsistent"),
    ("60-69", "0.6-0.69", "Clear, structured, mostly correct"),
    ("70-84", "0.7-0.84", "Polished, coherent, strong explanation"),
    ("85+", "0.85-1.0", "Publication-quality clarity"),
]

CRITERIA = [
    {
        "id": "document_structure",
        "title": "Document Structure",
        "signals": [
            "Logical flow between sections",
            "Proper use of headings and report sections",
            "Clear transitions between ideas",
        ],
    },
    {
        "id": "language_clarity",
        "title": "Language Clarity",
        "signals": [
            "Grammar and sentence control",
            "Formal academic tone",
            "Readable, precise wording",
        ],
    },
    {
        "id": "technical_explanation",
        "title": "Technical Explanation",
        "signals": [
            "Correct explanation of the model",
            "Appropriate use of technical terminology",
            "Clear reasoning for design choices",
        ],
    },
    {
        "id": "analysis_depth",
        "title": "Analysis Depth",
        "signals": [
            "Insight into results",
            "Quality of error analysis",
            "Evidence of critical thinking",
        ],
    },
    {
        "id": "referencing_and_academic_integrity",
        "title": "Referencing and Academic Integrity",
        "signals": [
            "Proper citations and attribution",
            "No plagiarism or copied analysis",
            "Transparent GenAI use where relevant",
        ],
    },
]

EXAMPLES = [
    {
        "label": "Weak",
        "band": "40-49",
        "score": 45,
        "workText": (
            "This project is about sentiment analysis. I used a dataset and trained a model. "
            "The accuracy was okay. The model sometimes failed but overall it worked fine. "
            "The results are shown below. The method is logistic regression. Some errors happened but I think it is fine."
        ),
        "reason": (
            "Very vague methodology description, no clear structure between sections, limited technical explanation, "
            "minimal error analysis, and informal language throughout."
        ),
    },
    {
        "label": "Average",
        "band": "60-69",
        "score": 62,
        "workText": (
            "We implemented a logistic regression classifier using TF-IDF features. The dataset consisted of labeled "
            "movie reviews. The model achieved an accuracy of 78%. Some errors were observed in cases of sarcasm and "
            "ambiguous phrasing. While the model performs reasonably well, it struggles with context-dependent sentiment."
        ),
        "reason": (
            "The methodology is clear and the structure is serviceable, with some awareness of model limitations, "
            "but the error analysis is still shallow and the transitions between ideas remain weak."
        ),
    },
    {
        "label": "Strong",
        "band": "70-84",
        "score": 78,
        "workText": (
            "We trained a logistic regression classifier using TF-IDF representations of the text. While the model "
            "achieved an accuracy of 81%, deeper analysis reveals systematic failure modes. In particular, the "
            "classifier struggles with negation (e.g., “not bad”) and sarcasm, which highlights its reliance on "
            "surface-level lexical cues rather than contextual understanding. This limitation suggests that more "
            "expressive models, such as transformer-based architectures, may better capture semantic nuances."
        ),
        "reason": (
            "Strong technical explanation, insightful error analysis, clear formal language, and a logical progression "
            "of ideas, though the discussion could still use more quantitative error breakdown and dataset bias analysis."
        ),
    },
    {
        "label": "Excellent",
        "band": "85+",
        "score": 90,
        "workText": (
            "Although the logistic regression model achieved an overall F1-score of 0.83, this aggregate metric obscures "
            "systematic weaknesses. A detailed error analysis reveals that misclassifications are disproportionately "
            "concentrated in reviews containing implicit sentiment, sarcasm, or domain-specific language. For instance, "
            "phrases such as “I expected more” were frequently misclassified due to their lack of explicit polarity markers. "
            "This suggests that the model’s reliance on bag-of-words representations limits its ability to capture "
            "compositional semantics. Future work should explore contextual embeddings (e.g., BERT) to address these limitations."
        ),
        "reason": (
            "Deep analytical insight, precise technical language, strong linkage between results and limitations, "
            "a professional academic tone, and well-justified next-step improvements."
        ),
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str | None = None, size: int = 10) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def apply_body_style(paragraph, *, bold: bool = False, size: int = 10, color: str = "1F2328") -> None:
    for run in paragraph.runs:
        run.font.name = "Aptos"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)


def build_rubric_markdown() -> str:
    lines = [
        f"# {RUBRIC_TITLE}",
        "",
        "## Coursework Task",
        f"**{TASK_OVERVIEW['title']}**",
        "",
        TASK_OVERVIEW["summary"],
        "",
        "### Requirements",
    ]
    lines.extend([f"- {item}" for item in TASK_OVERVIEW["requirements"]])
    lines.extend(
        [
            "",
            "### Deliverable",
            TASK_OVERVIEW["deliverable"],
            "",
            "## Rubric Focus",
            "This rubric scores the communication of ideas in the report.",
            "",
            "### Atomic Criteria",
        ]
    )
    lines.extend([f"- `{criterion['id']}`" for criterion in CRITERIA])
    lines.extend(
        [
            "",
            "## Band Mapping",
            "| Band | Score Range | Description |",
            "| --- | --- | --- |",
        ]
    )
    lines.extend([f"| {band} | {score_range} | {description} |" for band, score_range, description in BAND_MAPPING])
    lines.append("")
    lines.append("## Criterion Definitions")
    for criterion in CRITERIA:
        lines.append(f"### {criterion['title']} (`{criterion['id']}`)")
        lines.extend([f"- {signal}" for signal in criterion["signals"]])
        lines.append("")
    lines.extend(
        [
            "## Teacher Guidance",
            "- Use the rubric text as the source of truth when grading.",
            "- Score the new submission against the communication criteria, not generic writing advice.",
            "- Use the example calibrations to stay aligned to the intended band boundaries.",
            "",
            "## Calibration Examples Summary",
        ]
    )
    for example in EXAMPLES:
        lines.extend(
            [
                f"### {example['label']} Example",
                f"- Score: {example['score']}/100",
                f"- Band: {example['band']}",
                f"- Reason: {example['reason']}",
                "",
            ]
        )
    return "\n".join(lines).strip() + "\n"


def build_teacher_config(markdown_text: str) -> dict:
    return {
        "rubricFileName": RUBRIC_FILE_NAME,
        "rubricText": markdown_text,
        "examples": [
            {
                "workText": example["workText"],
                "score": example["score"],
                "reason": example["reason"],
            }
            for example in EXAMPLES
        ],
    }


def build_docx() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(RUBRIC_TITLE)
    title_run.bold = True
    title_run.font.size = Pt(21)
    title_run.font.color.rgb = RGBColor.from_string("243447")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Teacher-mode source document for the quick feedback prototype")
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor.from_string("5F645F")

    doc.add_paragraph("")

    heading = doc.add_paragraph("Coursework Task")
    apply_body_style(heading, bold=True, size=14, color="243447")

    body = doc.add_paragraph()
    body.add_run(TASK_OVERVIEW["title"] + ": ").bold = True
    body.add_run(TASK_OVERVIEW["summary"])
    apply_body_style(body, size=10)

    requirements_heading = doc.add_paragraph("Requirements")
    apply_body_style(requirements_heading, bold=True, size=12, color="243447")
    for item in TASK_OVERVIEW["requirements"]:
        paragraph = doc.add_paragraph(style=None)
        paragraph.style = doc.styles["List Bullet"]
        paragraph.add_run(item)
        apply_body_style(paragraph, size=10)

    deliverable = doc.add_paragraph()
    deliverable.add_run("Deliverable: ").bold = True
    deliverable.add_run(TASK_OVERVIEW["deliverable"])
    apply_body_style(deliverable, size=10)

    rubric_heading = doc.add_paragraph("Rubric Focus")
    apply_body_style(rubric_heading, bold=True, size=14, color="243447")

    rubric_intro = doc.add_paragraph(
        "This rubric scores communication of ideas in the report. It is intended to calibrate the LLM against teacher expectations."
    )
    apply_body_style(rubric_intro, size=10)

    criteria_heading = doc.add_paragraph("Atomic Criteria")
    apply_body_style(criteria_heading, bold=True, size=12, color="243447")
    for criterion in CRITERIA:
        paragraph = doc.add_paragraph(style=None)
        paragraph.style = doc.styles["List Bullet"]
        run = paragraph.add_run(f"{criterion['title']} ({criterion['id']})")
        run.bold = True
        apply_body_style(paragraph, size=10)

    band_heading = doc.add_paragraph("Band Mapping")
    apply_body_style(band_heading, bold=True, size=12, color="243447")

    band_table = doc.add_table(rows=1, cols=3)
    band_table.style = "Table Grid"
    band_table.autofit = False
    band_widths = [Inches(1.05), Inches(1.25), Inches(4.85)]
    headers = ["Band", "Score Range", "Description"]
    for idx, header in enumerate(headers):
        cell = band_table.rows[0].cells[idx]
        cell.width = band_widths[idx]
        set_cell_shading(cell, "243447")
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=10)
    for band, score_range, description in BAND_MAPPING:
        row = band_table.add_row()
        values = [band, score_range, description]
        for idx, value in enumerate(values):
            row.cells[idx].width = band_widths[idx]
            set_cell_text(row.cells[idx], value, size=10)

    criteria_detail_heading = doc.add_paragraph("Criterion Definitions")
    apply_body_style(criteria_detail_heading, bold=True, size=12, color="243447")
    for criterion in CRITERIA:
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(criterion["title"])
        title_run.bold = True
        apply_body_style(title_p, size=10)
        for signal in criterion["signals"]:
            paragraph = doc.add_paragraph(style=None)
            paragraph.style = doc.styles["List Bullet 2"]
            paragraph.add_run(signal)
            apply_body_style(paragraph, size=10)

    examples_heading = doc.add_paragraph("Calibration Examples")
    apply_body_style(examples_heading, bold=True, size=14, color="243447")
    intro = doc.add_paragraph(
        "These examples should be used as prompt-time anchors so the grader stays aligned with the teacher's interpretation of the rubric."
    )
    apply_body_style(intro, size=10)

    for example in EXAMPLES:
        card_title = doc.add_paragraph()
        card_title.add_run(f"{example['label']} Example").bold = True
        apply_body_style(card_title, size=12, color="243447")

        meta = doc.add_paragraph()
        meta.add_run(f"Score: {example['score']}/100   |   Band: {example['band']}")
        apply_body_style(meta, size=10, color="BF5C42")

        excerpt = doc.add_paragraph()
        excerpt.add_run("Text excerpt: ").bold = True
        excerpt.add_run(example["workText"])
        apply_body_style(excerpt, size=10)

        rationale = doc.add_paragraph()
        rationale.add_run("Teacher rationale: ").bold = True
        rationale.add_run(example["reason"])
        apply_body_style(rationale, size=10)

    guidance_heading = doc.add_paragraph("Quick-Feedback Input Guidance")
    apply_body_style(guidance_heading, bold=True, size=12, color="243447")
    guidance = [
        "Use the rubric text as the grading source of truth.",
        "Use the examples to calibrate scoring style, not to replace the rubric.",
        "Return criterion-level feedback, an overall score, and specific missing elements.",
    ]
    for item in guidance:
        paragraph = doc.add_paragraph(style=None)
        paragraph.style = doc.styles["List Bullet"]
        paragraph.add_run(item)
        apply_body_style(paragraph, size=10)

    return doc


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    markdown_text = build_rubric_markdown()
    (ROOT / RUBRIC_TEXT_FILE_NAME).write_text(markdown_text, encoding="utf-8")

    teacher_config = build_teacher_config(markdown_text)
    (ROOT / TEACHER_CONFIG_FILE_NAME).write_text(json.dumps(teacher_config, indent=2), encoding="utf-8")
    (ROOT / TEACHER_BUNDLE_FILE_NAME).write_text(json.dumps(teacher_config, indent=2), encoding="utf-8")

    doc = build_docx()
    doc.save(ROOT / RUBRIC_FILE_NAME)


if __name__ == "__main__":
    main()
